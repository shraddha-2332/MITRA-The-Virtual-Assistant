import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_callout(doc, text, bold_prefix=""):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F0F6F3")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left border only
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="146548"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_bold = p.add_run(bold_prefix + " ")
        r_bold.bold = True
        r_bold.font.color.rgb = RGBColor(0x14, 0x65, 0x48)
        r_bold.font.name = "Calibri"
        r_bold.font.size = Pt(10.5)
    r_text = p.add_run(text)
    r_text.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
    r_text.font.name = "Calibri"
    r_text.font.size = Pt(10)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def format_row(row, heights=None, bg_color=None, is_header=False):
    for cell in row.cells:
        if bg_color:
            set_cell_background(cell, bg_color)
        set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.1
            for r in p.runs:
                r.font.name = "Calibri"
                if is_header:
                    r.bold = True
                    r.font.size = Pt(10)
                    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                else:
                    r.font.size = Pt(9.5)
                    r.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

def build_handover_docx():
    doc = Document()
    
    # Page setup - 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header / Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.text = "MITRA: Virtual Assistant — Developer Architecture & Technical Handover Document"
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hp.runs[0].font.name = "Calibri"
        hp.runs[0].font.size = Pt(8.5)
        hp.runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.text = "MIT College of Railway Engineering and Research (MITCORER), Barshi | Confidential — Internal Technical Documentation"
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.runs[0].font.name = "Calibri"
        fp.runs[0].font.size = Pt(8.5)
        fp.runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    # -------------------------------------------------------------
    # Document Title Block
    # -------------------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    run_title = title_p.add_run("MITRA: THE VIRTUAL ASSISTANT")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(22)
    run_title.bold = True
    run_title.font.color.rgb = RGBColor(0x14, 0x65, 0x48) # Forest Green
    
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(12)
    run_sub = sub_p.add_run("Developer Architecture, Technical Briefing & System Handover Guide")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(14)
    run_sub.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)

    # Metadata Box Table
    meta_tbl = doc.add_table(rows=5, cols=2)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Target System", "MITRA (Multi-turn Institutional Transit & Regulatory Assistant)"),
        ("Target Institution", "MIT College of Railway Engineering and Research (MITCORER), Barshi (DTE: 06901)"),
        ("Version & Release Date", "Release v1.2 (October 2026) — Production Ready Prototype"),
        ("Document Purpose", "Developer Onboarding, System Architecture Briefing & Mail Handover Attachment"),
        ("Core Engineering Paradigm", "Deterministic Canonical Calculation + Constrained Decoding Models (CDM) + Hybrid Vector RAG"),
    ]
    for i, (k, v) in enumerate(meta_data):
        row = meta_tbl.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.text = k
        c1.text = v
        c0.paragraphs[0].runs[0].bold = True
        c0.paragraphs[0].runs[0].font.size = Pt(9.5)
        c0.paragraphs[0].runs[0].font.name = "Calibri"
        c0.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x14, 0x65, 0x48)
        c1.paragraphs[0].runs[0].font.size = Pt(9.5)
        c1.paragraphs[0].runs[0].font.name = "Calibri"
        set_cell_background(c0, "EBF3EE" if i % 2 == 0 else "F4F8F5")
        set_cell_background(c1, "FBFDFB" if i % 2 == 0 else "FFFFFF")
        set_cell_margins(c0, top=60, bottom=60, left=100, right=100)
        set_cell_margins(c1, top=60, bottom=60, left=100, right=100)
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # -------------------------------------------------------------
    # 1. EXECUTIVE SUMMARY & PROBLEM STATEMENT
    # -------------------------------------------------------------
    h1 = doc.add_heading("1. Executive Summary & Problem Context", level=1)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(4)
    h1.runs[0].font.color.rgb = RGBColor(0x14, 0x65, 0x48)

    doc.add_paragraph(
        "During academic admission cycles, candidates and guardians repeatedly ask high-stakes questions regarding "
        "branch options, sanctioned seat matrices, eligibility rules (PCM aggregate thresholds), category-wise fee concessions "
        "(Open, EBC, OBC, SC, ST, VJNT, TFWS), hostel deposits, bus transport, and scholarship portals. "
        "Traditional institutional websites present this data across scattered PDFs and notices that are cumbersome to navigate, "
        "while generic commercial LLMs (e.g. vanilla ChatGPT or Gemini) frequently hallucinate—producing incorrect fees, cutoffs, "
        "or false admission guarantees that present severe legal and financial liabilities."
    )
    
    create_callout(
        doc,
        "MITRA completely eliminates LLM hallucinations on critical institutional facts by never allowing a generative "
        "language model to invent statutory figures. All fees, intake quotas, DTE choice codes, and hostel charges are rendered "
        "deterministically from verified canonical records. Generative AI is used strictly for contextual language synthesis on "
        "general topics under a 15-second timeout and deterministic fallback.",
        "Core Architectural Rule:"
    )

    # -------------------------------------------------------------
    # 2. WHAT EXACTLY THE CHATBOT COVERS (FUNCTIONAL SCOPE)
    # -------------------------------------------------------------
    h2 = doc.add_heading("2. What Exactly the System Covers (Functional Scope)", level=1)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(4)
    h2.runs[0].font.color.rgb = RGBColor(0x14, 0x65, 0x48)

    doc.add_paragraph("MITRA is structured into two distinct execution paths: the Public Assistant Path and the Protected Maintainer Path.")

    doc.add_heading("2.1 Public Student & Parent Inquiries (Stateless Chat)", level=2)
    public_points = [
        ("First Year & DSE Fee Structures: ", "Accurate category-specific fee schedules (Open/OMS: ₹1,05,000; OBC/EBC Boys: ₹61,109; Girls/VJNT/TFWS: ₹17,217; SC/ST: ₹8,000) for AY 2026-27 with DSE lateral entry breakdowns and official circular snapshots."),
        ("B.Tech Programs & Sanctioned Intakes: ", "Covers Computer Science & Engineering (90 seats), Artificial Intelligence & Data Science (60 seats), Mechanical & Rail Engineering (60 seats), Civil Engineering (30 seats), Electronics & Telecommunication (60 seats), and Electronics & Computer Engineering (60 seats)."),
        ("MHT-CET & JEE Main Cutoffs: ", "Historical CAP Round 1, 2, and 3 cutoff percentiles and merit trends for institute choice codes (DTE Code: 06901)."),
        ("Residential & Campus Facilities: ", "Hostel accommodation (₹35,000/yr + ₹5,000 refundable deposit), annual mess fees (₹42,000/yr), railway model lab, Central Workshop machine sections, bus transport routes, and library digital resources."),
        ("Student Associations & Field Training: ", "Departmental bodies including SPACE (Students' Association of Computer Engineering), MESA (Mechanical), CESA (Civil), Rail-MIT Symposium, and practical visits to Kurduwadi Railway Workshop and Solapur Divisional HQ."),
        ("Multi-Turn Follow-Ups: ", "Resolves coreferences and ellipsis across conversation history (e.g., 'What about OBC girls?' after asking about CSE fees) without losing context."),
        ("Trilingual Support (English, Marathi, Hindi): ", "Processes Devanagari script, Marathi vernacular, Hindi inquiries, and Roman-script transliteration (e.g. 'Hostel chi fee kiti ahe?')."),
        ("Human Escalation Desk: ", "Out-of-distribution queries (e.g. staff hiring, faculty recruitment, unverified policies) trigger an escalation card providing official admission phone numbers and capturing candidate inquiries for staff follow-up.")
    ]
    for bold_prefix, text in public_points:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        r_b = p.add_run(bold_prefix)
        r_b.bold = True
        p.add_run(text)

    doc.add_heading("2.2 Admin & Maintenance Portal (Role-Protected)", level=2)
    admin_points = [
        ("Official Source Document Ingestion: ", "Upload institutional notices, government gazettes, and cutoff circulars in PDF, DOCX, XLSX, CSV, TXT, and image formats."),
        ("Constrained Decoding Model (CDM) Extraction: ", "One-click automated extraction using Gemini's responseSchema to enforce strict JSON fact vectors (category, title, text, tags) with zero syntax errors."),
        ("Human-in-the-Loop Review Drawer: ", "Admin can inspect, edit, and approve extracted text before promoting it from 'Review' to 'Published' status."),
        ("Automated Website Drift Synchronization: ", "Periodically crawls the official college domain (mitcorer.edu.in), detects text/year changes, and stages modified pages for review without dangerous auto-publishing."),
        ("Concurrency Control: ", "Revision-locked editing prevents two administrators from accidentally overwriting concurrent changes.")
    ]
    for bold_prefix, text in admin_points:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        r_b = p.add_run(bold_prefix)
        r_b.bold = True
        p.add_run(text)

    # -------------------------------------------------------------
    # 3. TECHNOLOGY STACK
    # -------------------------------------------------------------
    h3 = doc.add_heading("3. Complete Technology Stack & Architecture", level=1)
    h3.paragraph_format.space_before = Pt(12)
    h3.paragraph_format.space_after = Pt(4)
    h3.runs[0].font.color.rgb = RGBColor(0x14, 0x65, 0x48)

    tech_table = doc.add_table(rows=8, cols=3)
    tech_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Layer / Subsystem", "Technologies Used", "Architectural Role & Significance"]
    tech_data = [
        ("Frontend Client", "React 19, Vite 8, Lucide-React, Pure CSS", "High-performance SPA, mobile-first floating chat widget, quick-reply chips, and admin portal. Zero heavy UI frameworks."),
        ("Backend Server", "Node.js (v20+ LTS), Built-in HTTP module", "Modular monolith API. Handles request routing, CORS, rate limiting, and JWT authentication without framework bloat."),
        ("Database & Storage", "SQLite (via native node:sqlite)", "Zero-setup embedded database storing canonical facts, official source files, revision histories, and vector embeddings."),
        ("AI & Generative Layer", "Google Gemini 2.5 Flash / 1.5 Flash", "Used only for open-ended contextual synthesis and CDM structured document parsing via responseSchema. 15s timeout with fallback."),
        ("Constrained Decoding (CDM)", "Gemini responseSchema + Deterministic Parser", "Guarantees that uploaded circulars are parsed into strict JSON objects conforming to the institutional fact schema."),
        ("Multilingual Translation", "Hash-Bound Reviewed Translations + English Fallback", "Provides verified Marathi and Hindi answers; automatically drops to English if upstream English fact revisions change."),
        ("Testing & Quality Assurance", "Node Test Runner (node:test), Playwright Core, Oxlint", "146 backend tests, 15 frontend tests, and 426 real-API browser journeys passing. Oxlint provides instantaneous linter verification.")
    ]
    for j, h in enumerate(headers):
        tech_table.rows[0].cells[j].text = h
    format_row(tech_table.rows[0], bg_color="146548", is_header=True)
    for i, row_data in enumerate(tech_data):
        row = tech_table.rows[i + 1]
        for j, val in enumerate(row_data):
            row.cells[j].text = val
        format_row(row, bg_color="F9FBFA" if i % 2 == 0 else "FFFFFF")
    tech_table.columns[0].width = Inches(1.8)
    tech_table.columns[1].width = Inches(2.2)
    tech_table.columns[2].width = Inches(2.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # 4. SYSTEM WORKFLOW & REQUEST LIFECYCLE
    # -------------------------------------------------------------
    h4 = doc.add_heading("4. Architectural Workflow & Request Lifecycle", level=1)
    h4.paragraph_format.space_before = Pt(12)
    h4.paragraph_format.space_after = Pt(4)
    h4.runs[0].font.color.rgb = RGBColor(0x14, 0x65, 0x48)

    doc.add_paragraph(
        "When a user submits a question to POST /api/public/assistant, the backend executes the following multi-stage pipeline:"
    )

    steps = [
        ("1. Input Screening & Rate Limiting: ", "Validates query string length and IP-based rate limiting to prevent denial-of-service."),
        ("2. Multi-Turn Context Condensation: ", "Evaluates the user's session history (last 4 turns). If the query is an elliptical follow-up (e.g. 'What about girls?'), the condenser injects previous parameters (branch: Computer Engineering, category: OBC) to form a self-contained query vector."),
        ("3. Intent & Dimension Extraction: ", "A bounded, dependency-free parser extracts the primary topic, candidate category, gender, academic year, and facility."),
        ("4. Deterministic Canonical Routing (Zero-Hallucination Gate): ", "If the topic involves fees, intakes, hostel charges, or DTE choice codes, the system directly invokes the Deterministic Renderer. Factual values are pulled from canonical records and formatted into verified tables. The LLM is bypassed completely."),
        ("5. Hybrid Retrieval for Open Queries: ", "For general inquiries, the system queries Published records via lexical keyword weighting and SQLite vector cosine similarity."),
        ("6. Evidence-Grounded Gemini Generation: ", "The retrieved evidence is injected into an evidence-constrained prompt. The candidate answer is passed through factual-token checks and anti-promissory validation before delivery."),
        ("7. Provenance & Citations: ", "The final answer is delivered with official source citations, academic year tags, fee circular images, and human helpline contacts.")
    ]
    for b_pre, b_text in steps:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(b_pre)
        r.bold = True
        p.add_run(b_text)

    # -------------------------------------------------------------
    # 5. KEY API ENDPOINTS FOR DEVELOPERS
    # -------------------------------------------------------------
    h5 = doc.add_heading("5. Key API Endpoints & Developer Contracts", level=1)
    h5.paragraph_format.space_before = Pt(12)
    h5.paragraph_format.space_after = Pt(4)
    h5.runs[0].font.color.rgb = RGBColor(0x14, 0x65, 0x48)

    api_table = doc.add_table(rows=8, cols=3)
    api_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    api_headers = ["Endpoint & Method", "Access Level", "Description & Request Payload"]
    api_data = [
        ("POST /api/public/assistant", "Public (No Auth)", "Main chat query. Accepts { question, language, conversationId }. Returns structured answer with source links."),
        ("GET /api/public/facts", "Public (No Auth)", "Fetches single-source-of-truth canonical facts: category fees, intake list, and hostel charges."),
        ("POST /api/public/inquiries", "Public (No Auth)", "Captures escalated candidate inquiries { name, contact, question, category } for human staff followup."),
        ("POST /api/auth/login", "Public / Maintainer", "Admin authentication. Body: { email, password }. Returns JWT token and sets httpOnly refresh cookie."),
        ("POST /api/official-sources/upload", "Admin Only (JWT)", "Uploads official source file (multipart/form-data) with title, documentType, academicYear."),
        ("POST /api/official-sources/:id/extract-structured", "Admin Only (JWT)", "Triggers Constrained Decoding Model (CDM) extraction. Returns validated keyFacts array adhering to JSON schema."),
        ("PATCH /api/official-sources/:id/publish", "Admin Only (JWT)", "Promotes reviewed text into live retrieval knowledge base. Body: { verifiedText }.")
    ]
    for j, h in enumerate(api_headers):
        api_table.rows[0].cells[j].text = h
    format_row(api_table.rows[0], bg_color="146548", is_header=True)
    for i, row_data in enumerate(api_data):
        row = api_table.rows[i + 1]
        for j, val in enumerate(row_data):
            row.cells[j].text = val
        format_row(row, bg_color="F9FBFA" if i % 2 == 0 else "FFFFFF")
    api_table.columns[0].width = Inches(2.2)
    api_table.columns[1].width = Inches(1.5)
    api_table.columns[2].width = Inches(2.8)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # 6. HOW TO RUN, DEVELOP, AND TEST
    # -------------------------------------------------------------
    h6 = doc.add_heading("6. Setup, Local Run & Testing Instructions", level=1)
    h6.paragraph_format.space_before = Pt(12)
    h6.paragraph_format.space_after = Pt(4)
    h6.runs[0].font.color.rgb = RGBColor(0x14, 0x65, 0x48)

    doc.add_paragraph("Follow these exact steps to run the complete environment locally:")

    setup_steps = [
        ("Step 1: Environment Configuration", "In backend/.env, set TOKEN_SECRET=your_jwt_secret, GEMINI_API_KEY=your_key, and GEMINI_MODEL=gemini-2.5-flash. If no API key is provided, the system seamlessly operates on deterministic fallback."),
        ("Step 2: Start Backend Server", "Run: npm --prefix backend run dev. The server listens on http://127.0.0.1:4000 and automatically initializes the SQLite schema and seed records."),
        ("Step 3: Start Frontend Dev Server", "Run: npm --prefix frontend run dev. The Vite dev server listens on http://127.0.0.1:5173."),
        ("Step 4: Accessing the Admin Portal", "Open http://127.0.0.1:5173/?admin in your browser (or click the '©' symbol in the footer). Login with Default Admin Email: admin@mitcorer.edu.in | Password: mitra."),
        ("Step 5: Executing Test Suites", "Run backend unit tests: npm --prefix backend test (146 tests passing). Run frontend unit tests: npm --prefix frontend test (15 tests passing). Run end-to-end browser tests: npm --prefix frontend run test:maintainer-browser.")
    ]
    for title, desc in setup_steps:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(title + ": ")
        r.bold = True
        r.font.color.rgb = RGBColor(0x14, 0x65, 0x48)
        p.add_run(desc)

    # -------------------------------------------------------------
    # 7. CODEBASE DIRECTORY MAP
    # -------------------------------------------------------------
    h7 = doc.add_heading("7. Codebase Directory Map for Developers", level=1)
    h7.paragraph_format.space_before = Pt(12)
    h7.paragraph_format.space_after = Pt(4)
    h7.runs[0].font.color.rgb = RGBColor(0x14, 0x65, 0x48)

    dir_table = doc.add_table(rows=8, cols=2)
    dir_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    dir_headers = ["File / Directory Path", "Developer Significance"]
    dir_data = [
        ("backend/src/server.js", "Central API server. Contains routes, bounded query parser, deterministic renderers, and the extractStructuredFactsWithCdm pipeline."),
        ("backend/src/seed.js", "Canonical ground-truth records (KB-1 to KB-51). Defines official fee slabs, intakes, hostel charges, and SPACE student body."),
        ("backend/src/website-sync.js", "Website monitoring crawler. Fetches live mitcorer.edu.in pages and detects numerical/year drift for admin staging."),
        ("frontend/src/App.jsx", "Primary React client. Contains PublicAssistantWidget, prompt chips, KnowledgeManager, and the admin login modal."),
        ("frontend/src/api.js", "Frontend REST client. Houses methods for asking questions, admin login, document upload, and CDM structured extraction."),
        ("backend/test/api.test.js", "146 comprehensive backend unit tests validating security, CDM schemas, fee calculations, and multi-turn flows."),
        ("frontend/scripts/verify-maintainer.mjs", "Playwright end-to-end browser test validating maintainer session restoration and knowledge editing.")
    ]
    for j, h in enumerate(dir_headers):
        dir_table.rows[0].cells[j].text = h
    format_row(dir_table.rows[0], bg_color="146548", is_header=True)
    for i, row_data in enumerate(dir_data):
        row = dir_table.rows[i + 1]
        for j, val in enumerate(row_data):
            row.cells[j].text = val
        format_row(row, bg_color="F9FBFA" if i % 2 == 0 else "FFFFFF")
    dir_table.columns[0].width = Inches(2.6)
    dir_table.columns[1].width = Inches(3.9)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # 8. NON-NEGOTIABLE ENGINEERING RULES
    # -------------------------------------------------------------
    h8 = doc.add_heading("8. Non-Negotiable Engineering Rules for Future Development", level=1)
    h8.paragraph_format.space_before = Pt(12)
    h8.paragraph_format.space_after = Pt(4)
    h8.runs[0].font.color.rgb = RGBColor(0x14, 0x65, 0x48)

    rules = [
        ("RULE 1: Never Let an LLM Calculate or Guess Statutory Numbers", "Fees, intakes, and cutoffs must always be pulled from canonical records in backend/src/seed.js or database published records. Never prompt an LLM to generate fee figures without deterministic grounding."),
        ("RULE 2: Publication Gating is Mandatory", "Records in Draft, Review, Archived, or Retired states must never be exposed to public retrieval. Only records in Published state are accessible."),
        ("RULE 3: Zero Auto-Publishing on Crawled Content", "The website synchronizer must only stage detected drift in the Review status. An authenticated administrator must manually review and publish updates."),
        ("RULE 4: Preserve Translation Fallbacks", "If an English canonical record changes, any translated output whose content hash does not match must explicitly fall back to verified English with a notice."),
        ("RULE 5: Zero-Dependency Philosophy", "Do not add heavy external vector databases, cloud dependencies, or ORMs. The lightweight Node.js + SQLite stack was deliberately chosen for low resource footprints and zero vendor lock-in.")
    ]
    for r_title, r_desc in rules:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(r_title + ": ")
        r.bold = True
        r.font.color.rgb = RGBColor(0x8B, 0x00, 0x00) # Dark red for rules
        p.add_run(r_desc)

    # Output path
    output_filename = "MITRA_Developer_Handover_Guide.docx"
    output_path = os.path.abspath(output_filename)
    doc.save(output_path)
    print(f"Successfully generated: {output_path}")

if __name__ == "__main__":
    build_handover_docx()

